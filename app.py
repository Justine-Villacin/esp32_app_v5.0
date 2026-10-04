"""
app.py — SmartGrow Flask web app.

This file contains ONLY Flask routes and WSGI-safe imports. Vercel imports
this module directly and calls the `app` object — it never runs the
`if __name__ == "__main__":` block, so anything that needs a real desktop
window, a background thread, or native GUI libraries must NOT live here.

For the pywebview desktop app, the manual HTTPS runner, and the dev-mode
file watcher, see desktop.py (local use only — `python desktop.py`).
"""

import csv
import io
import os
import queue
import threading
import time

from flask import Flask, Response, jsonify, render_template, request, send_file

import storage
from openpyxl import Workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Static assets (css/js/icons/manifest/sw.js) live in public/ at the project
# root. Vercel's CDN serves that folder directly and ignores Flask's static
# handling entirely, but pointing Flask's own static_folder at the same
# folder means `python desktop.py` (plain Flask dev server, no Vercel CDK in
# front of it) serves the exact same files at the exact same URLs — no
# "/static" prefix, no duplicate routes needed for manifest.json / sw.js.
app = Flask(__name__, static_folder=os.path.join(BASE_DIR, "public"), static_url_path="")

# ============================================================
# DEV MODE — live reload (local development only)
# ============================================================
# Defaults to OFF. desktop.py sets SMARTGROW_DEV_MODE=true before importing
# this module so local runs keep the old "always on" dev experience; a
# deployment on Vercel never sets this, so it's off unless someone opts in
# on purpose. When it's off, the /dev/reload-stream route below is never
# registered at all (not just short-circuited) — hitting it 404s exactly
# like any other unknown route, and nothing is left running or listening.
DEV_MODE = os.environ.get("SMARTGROW_DEV_MODE", "false").strip().lower() in ("1", "true", "yes", "on")

# Optional shared-secret required on /update and /api/control once this
# dashboard is reachable from the public internet. If it's unset, those
# routes stay open (matching the original home-LAN-only behavior) — set it
# in your environment before pointing a public *.vercel.app URL at real
# hardware. See README.md / .env.example.
API_KEY = os.environ.get("SMARTGROW_API_KEY")

_reload_listeners = []
_reload_lock = threading.Lock()


def broadcast_reload():
    """Called by desktop.py's file watcher when a template/static file
    changes, so every connected browser tab reloads itself."""
    with _reload_lock:
        for q in _reload_listeners:
            q.put("reload")


if DEV_MODE:
    @app.route("/dev/reload-stream")
    def dev_reload_stream():
        """Server-Sent-Events stream for the live-reload client in
        index.html. Only registered at all when DEV_MODE is on."""
        def event_stream():
            q = queue.Queue()
            with _reload_lock:
                _reload_listeners.append(q)
            try:
                while True:
                    msg = q.get()
                    yield f"data: {msg}\n\n"
            finally:
                with _reload_lock:
                    if q in _reload_listeners:
                        _reload_listeners.remove(q)

        return Response(event_stream(), mimetype="text/event-stream")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _check_api_key(req):
    """Returns True if the request is allowed to proceed. When API_KEY isn't
    configured, every request is allowed (dev/home-LAN mode)."""
    if not API_KEY:
        return True
    return req.headers.get("X-API-Key") == API_KEY


def _public_state(state):
    """Strips internal bookkeeping fields before a dict is sent to the client."""
    return {k: v for k, v in state.items() if not k.startswith("_")}


def _validate_reading(data):
    """Validates/coerces the ESP32 payload for /update. Returns
    (clean_values, errors). `clean_values` only contains keys that passed
    validation; on any error, the whole request is rejected by the caller."""
    errors = []
    values = {}

    def bounded_float(key, lo, hi):
        if key not in data:
            return
        raw = data.get(key)
        try:
            val = float(raw)
        except (TypeError, ValueError):
            errors.append(f"'{key}' must be a number")
            return
        if not (lo <= val <= hi):
            errors.append(f"'{key}' must be between {lo} and {hi}")
            return
        values[key] = val

    bounded_float("temperature", -40.0, 80.0)
    bounded_float("humidity", 0.0, 100.0)
    bounded_float("soil_moisture", 0.0, 100.0)
    # "light" is the numeric LDR reading (0-1000, see firmware/config.h).
    # NOTE: this is intentionally a different JSON key from the grow-light
    # RELAY state ("grow_light", "ON"/"OFF") — the original firmware/API
    # contract used "light" for both, which silently broke the numeric
    # reading. See CHANGELOG.md.
    bounded_float("light", 0.0, 1000.0)

    return values, errors


def _export_unavailable(message, status=200):
    """A styled, frontend-friendly JSON response for the CSV/Excel export
    routes, used instead of a bare text 404. The frontend's downloadFile()
    helper (public/js/utils.js) catches this and shows it as a toast in the
    activity feed instead of a jarring unstyled page."""
    return jsonify({"status": "unavailable", "message": message}), status


def _local_csv_path():
    """Path to the desktop-mode CSV log. Only ever exists when the app is
    running locally via desktop.py — never on Vercel's read-only/ephemeral
    filesystem."""
    return os.path.join(BASE_DIR, "sensor_data.csv")


# ---------------------------------------------------------------------------
# Page + status routes
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template("index.html", dev_mode=DEV_MODE)


@app.route("/api/sensors")
def api_sensors():
    state, connected = storage.touch()
    out = _public_state(state)
    if not connected:
        # The ESP32 hasn't been heard from recently: we genuinely don't know
        # the real state of the relays anymore (the last report could be
        # stale), so never keep showing them as ACTIVE. Combined with the
        # firmware's own fail-safes (see firmware/SmartGrow_ESP32.ino), this
        # closes both the server-reported and hardware-level versions of
        # "fan/pump/light stay on when disconnected."
        out["esp32Connected"] = False
        out["dhtConnected"] = False
        out["soilConnected"] = False
        out["lightConnected"] = False
        out["fan"] = False
        out["pump"] = False
        out["lightState"] = False
    return jsonify(out)


@app.route("/api/status")
def api_status():
    """Small diagnostic endpoint — confirms which storage backend is active
    (redis vs. in-memory fallback) and whether the ESP32 currently looks
    connected. Handy for verifying a Vercel + Redis deployment."""
    _, connected = storage.touch()
    return jsonify({
        "storage_backend": storage.backend_name(),
        "redis_url_set": storage.is_remote_configured(),
        "redis_error": storage.last_redis_error(),
        "dev_mode": DEV_MODE,
        "api_key_required": bool(API_KEY),
        "esp32_connected": connected,
        "disconnect_timeout_seconds": storage.DISCONNECT_TIMEOUT_SECONDS,
    })


@app.route("/api/toggle/<sensor_name>")
def toggle_sensor(sensor_name):
    key_map = {"dht": "dhtConnected", "soil": "soilConnected", "light": "lightConnected"}
    if sensor_name not in key_map:
        return jsonify({"status": "error", "message": f"unknown sensor '{sensor_name}'"}), 400

    state = storage.get_latest()
    key = key_map[sensor_name]
    state[key] = not state.get(key, False)
    storage.set_latest(state)
    return jsonify(_public_state(state))


# --- handles the manual pump/light/fan buttons in the dashboard UI ---
@app.route("/api/control", methods=["POST"])
def api_control():
    if not _check_api_key(request):
        return jsonify({"status": "error", "message": "unauthorized"}), 401

    data = request.get_json(silent=True)
    if data is None:
        return jsonify({"status": "error", "message": "invalid or missing JSON body"}), 400

    target = data.get("target")
    key_map = {"pump": "pump", "light": "lightState", "fan": "fan"}
    if target not in key_map:
        return jsonify({"status": "error", "message": f"unknown target '{target}'"}), 400

    state = storage.get_latest()
    state[key_map[target]] = bool(data.get("state"))
    storage.set_latest(state)
    return jsonify({"status": "success", "sensor_data": _public_state(state)})


# --- receives live JSON from the physical ESP32 ---
@app.route("/update", methods=["POST"])
def update_from_esp32():
    if not _check_api_key(request):
        return jsonify({"status": "error", "message": "unauthorized"}), 401

    data = request.get_json(silent=True)
    if data is None:
        return jsonify({"status": "error", "message": "invalid or missing JSON body"}), 400

    values, errors = _validate_reading(data)
    if errors:
        return jsonify({"status": "error", "message": "; ".join(errors)}), 400

    now = time.time()
    state = storage.get_latest()

    if "temperature" in values:
        state["temp"] = values["temperature"]
    if "humidity" in values:
        state["hum"] = values["humidity"]
    if "soil_moisture" in values:
        state["soil"] = values["soil_moisture"]
    if "light" in values:
        state["light"] = values["light"]

    state["esp32Connected"] = True
    # Convert "ON"/"OFF" strings from the ESP32 firmware to booleans.
    # "grow_light" is the RELAY command, deliberately a different JSON key
    # from "light" (the numeric LDR reading above) — see firmware/README.md
    # and CHANGELOG.md for why these used to collide.
    state["fan"] = (data.get("fan") == "ON")
    state["pump"] = (data.get("water_pump") == "ON")
    state["lightState"] = (data.get("grow_light") == "ON")

    # Real per-sensor status the ESP32 reports, falling back to True for
    # older firmware that doesn't send these fields yet.
    state["dhtConnected"] = bool(data.get("dht_connected", True))
    state["soilConnected"] = bool(data.get("soil_connected", True))
    state["lightConnected"] = bool(data.get("light_connected", True))
    state["_last_update_time"] = now

    # Defense in depth: the firmware is already supposed to force its own
    # relay off the instant a sensor read fails (see
    # firmware/SmartGrow_ESP32.ino's runAutomation()), but we don't want a
    # firmware bug, a stale/replayed payload, or a modified device to be
    # able to report "sensor disconnected" and "actuator ON" at the same
    # time and have the dashboard repeat that uncritically. Enforce the
    # same rule server-side too.
    if not state["dhtConnected"]:
        state["fan"] = False
    if not state["soilConnected"]:
        state["pump"] = False
    if not state["lightConnected"]:
        state["lightState"] = False

    storage.set_latest(state)
    storage.append_history({
        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(now)),
        "temp": state["temp"] if state["dhtConnected"] else "ERR",
        "hum": state["hum"] if state["dhtConnected"] else "ERR",
        "soil": state["soil"] if state["soilConnected"] else "ERR",
        "light": state["light"] if state["lightConnected"] else "ERR",
        "pump": state["pump"],
        "lightState": state["lightState"],
        "fan": state["fan"],
    })

    return jsonify({"status": "success"})


# ---------------------------------------------------------------------------
# Data export (CSV / Excel)
# ---------------------------------------------------------------------------
# History source of truth, in priority order:
#   1. The local desktop CSV file written by desktop.py's background_logging
#      thread — only ever present when running locally via desktop.py, and
#      persists across restarts, so it wins when it exists.
#   2. storage.get_history() — Redis-backed once REDIS_URL is configured
#      (this is option (a) from the brief), or the in-memory fallback
#      otherwise (handy for a quick `vercel dev`/local test with no Redis
#      attached yet, though it won't survive a real serverless cold start).
#   3. Neither has any data -> a styled "not available yet" response instead
#      of a bare 404 (option (b), used as the graceful fallback rather than
#      the primary strategy — see CHANGELOG.md for why).

CSV_HEADERS = [
    "Timestamp", "Temperature_C", "Humidity_pct", "Soil_Moisture_pct",
    "Light_Lux", "Pump_Active", "Grow_Light_Active", "Exhaust_Fan_Active",
]


def _history_rows_from_store():
    rows = storage.get_history()
    return [
        [r.get("timestamp"), r.get("temp"), r.get("hum"), r.get("soil"),
         r.get("light"), r.get("pump"), r.get("lightState"), r.get("fan")]
        for r in rows
    ]


def _history_rows_from_local_csv():
    path = _local_csv_path()
    if not os.path.exists(path):
        return None
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        next(reader, None)  # skip header, we use our own canonical CSV_HEADERS
        for row in reader:
            if row:
                rows.append(row)
    return rows


def _resolve_history_rows():
    """Returns (rows, source) where source is 'local', 'redis', 'memory',
    or None if there's genuinely nothing to export yet."""
    local_rows = _history_rows_from_local_csv()
    if local_rows is not None:
        return local_rows, "local"

    rows = _history_rows_from_store()
    if rows:
        return rows, storage.backend_name()

    if storage.is_remote_configured():
        # Redis is wired up but genuinely empty so far (e.g. a brand-new
        # deployment that hasn't heard from the ESP32 yet).
        return [], storage.backend_name()

    return None, None


@app.route("/sensor_data.csv")
def download_csv_file():
    rows, source = _resolve_history_rows()
    if source is None:
        return _export_unavailable(
            "No sensor history is available yet. This appears once your "
            "ESP32 has sent at least one reading (and, on a hosted "
            "deployment, once a Redis connection is configured — see README.md)."
        )

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(CSV_HEADERS)
    writer.writerows(rows)

    mem = io.BytesIO(buffer.getvalue().encode("utf-8"))
    return send_file(
        mem,
        as_attachment=True,
        download_name="SmartGrow_Complete_Data.csv",
        mimetype="text/csv",
    )


@app.route("/download/excel")
def download_excel_file():
    rows, source = _resolve_history_rows()
    if source is None:
        return _export_unavailable(
            "No sensor history is available yet, so there's nothing to "
            "export to Excel yet. This appears once your ESP32 has sent "
            "readings (and, on a hosted deployment, once a Redis connection "
            "is configured — see README.md)."
        )

    wb = Workbook()
    wb.remove(wb.active)  # remove the default empty sheet

    days_data = {}
    for row in rows:
        if not row or str(row[0]).startswith("---"):
            continue
        date_str = str(row[0]).split(" ")[0]
        days_data.setdefault(date_str, []).append(row)

    if days_data:
        for date_str, day_rows in days_data.items():
            ws = wb.create_sheet(title=date_str)
            ws.append(CSV_HEADERS)
            for r in day_rows:
                ws.append(r)
    else:
        ws = wb.create_sheet(title="Data")
        ws.append(CSV_HEADERS)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)

    return send_file(
        mem,
        as_attachment=True,
        download_name="SmartGrow_Data.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )